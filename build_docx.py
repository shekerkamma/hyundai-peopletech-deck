"""Build the Word version of the executive brief."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1A, 0x27, 0x54)
TEAL = RGBColor(0x14, 0xA0, 0x85)
INK = RGBColor(0x1E, 0x29, 0x3B)
INK_SOFT = RGBColor(0x47, 0x55, 0x69)
INK_MUTE = RGBColor(0x64, 0x74, 0x8B)

doc = Document()

# Tight margins for one-page memo
for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

def add_para(text, *, size=10.5, bold=False, italic=False, color=INK,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return p

def add_section_header(text, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = color
    # Add bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "14A085")
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_bullet(parts):
    """parts = [(text, bold), ...] for mixed-format bullet."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    for text, bold in parts:
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.bold = bold
        run.font.color.rgb = INK
    return p

def add_mixed_para(parts, space_after=4):
    """A paragraph with mixed bold/regular text."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    for text, bold in parts:
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.bold = bold
        run.font.color.rgb = INK
    return p

# ============================================================================
# HEADER
# ============================================================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(2)
r = title.add_run("EXECUTIVE BRIEF")
r.font.name = "Calibri"; r.font.size = Pt(10); r.font.bold = True
r.font.color.rgb = TEAL

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(6)
r = sub.add_run("AI for Plant Operations: A 12-Week Path to Measurable ROI")
r.font.name = "Calibri"; r.font.size = Pt(18); r.font.bold = True
r.font.color.rgb = NAVY

meta_text = (
    "To: Hyundai · Plant Directors · Operations Leaders · Quality Management   |   "
    "From: PeopleTech · AI Manufacturing Practice   |   May 2026"
)
mp = doc.add_paragraph()
mp.paragraph_format.space_after = Pt(2)
r = mp.add_run(meta_text)
r.font.name = "Calibri"; r.font.size = Pt(9); r.font.italic = True
r.font.color.rgb = INK_MUTE

mp2 = doc.add_paragraph()
mp2.paragraph_format.space_after = Pt(8)
r = mp2.add_run("Companion to the 15-slide deck \"AI for Plant Operations · Hyundai × PeopleTech\".")
r.font.name = "Calibri"; r.font.size = Pt(9); r.font.italic = True
r.font.color.rgb = INK_MUTE

# Divider via bottom border on an empty paragraph
divp = doc.add_paragraph()
divp.paragraph_format.space_before = Pt(0)
divp.paragraph_format.space_after = Pt(6)
pPr = divp._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "12")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "14A085")
pBdr.append(bottom)
pPr.append(pBdr)

# ============================================================================
# EXECUTIVE SUMMARY
# ============================================================================
add_section_header("Executive Summary")
add_para(
    "Automotive manufacturing has crossed an AI inflection point. Edge AI hardware is mature, "
    "hyperscaler vision and ML platforms are production-grade, and three years of peer-OEM deployments "
    "now provide a settled answer to \"does it work?\" — it does, with ROI ranging from 3:1 to over 5:1 "
    "within the first year."
)
add_mixed_para([
    ("The decision in front of Hyundai is no longer whether to deploy AI on the plant floor, but ", False),
    ("where to start", True), (", ", False),
    ("how fast to prove value", True), (", and ", False),
    ("which partner can execute against Hyundai's uptime and quality bars without disrupting the line", True),
    (". This brief recommends a ", False),
    ("12-week pilot anchored on two use cases", True),
    (" — AI-based Visual Inspection and Predictive Maintenance — that map directly to Hyundai's stated key areas, carry the highest peer-validated ROI, and create the data and integration scaffolding for the remaining six use cases detailed in the accompanying deck.", False),
])

# ============================================================================
# THE OPPORTUNITY
# ============================================================================
add_section_header("The Opportunity")
add_para("Across the eight use cases mapped to Hyundai's key areas, the value pool concentrates in five drivers — each backed by measured outcomes from comparable production deployments:")
add_bullet([("Quality escapes & warranty exposure. ", True),
            ("AI Visual Inspection cuts escape defects by up to 60% and rework cost by 45%. Across paint, weld, and assembly surfaces of a typical OEM line, this compounds into eight- to nine-figure annual avoided cost.", False)])
add_bullet([("Unplanned downtime. ", True),
            ("Predictive Maintenance reduces unplanned downtime by 40% on average, with 5:1 measured ROI in year one. At industry-standard downtime cost of $22,000/minute, even single-line gains pay back the program multiple times over.", False)])
add_bullet([("Variant & fitment errors. ", True),
            ("Variant Confirmation cuts wrong-variant assembly by 90%; Seating & Component Validation cuts fitment errors by 85% — eliminating an entire class of recall risk before vehicles leave the station.", False)])
add_bullet([("SOP & safety compliance. ", True),
            ("Vision + pose estimation drives SOP deviations down 65% and safety incidents down 55%, with audit-ready EHS reporting that materially reduces regulatory exposure.", False)])
add_bullet([("Traceability & recall response. ", True),
            ("Full digital genealogy cuts recall investigation from weeks to days — turning a compliance liability into a competitive advantage.", False)])
add_para(
    "These are not projections. They are outcome ranges from in-production AI deployments at Tier-1 OEMs, EV manufacturers, and adjacent automotive supply industries.",
    italic=True, color=INK_SOFT, size=10
)

# ============================================================================
# WHY PEOPLETECH
# ============================================================================
add_section_header("Why PeopleTech")
add_mixed_para([
    ("A complement, not a replacement, for Hyundai's AI bets. ", True),
    ("Hyundai Motor Group already operates HMGICS Singapore as a smart-factory testbed and has invested in AI through AIRS Company and Boston Dynamics. PeopleTech's role is to bring those innovations to production scale at high-volume plants — Ulsan, Asan, HMGMA Metaplant America, HMMA, HMMC, HMI — and operationalize AI on lines already running today.", False),
])
add_mixed_para([
    ("Manufacturing domain depth. ", True),
    ("Our engineers have shipped vision QA, predictive maintenance, and traceability stacks for Tier-1 OEMs, EV manufacturers, aerospace suppliers, and heavy equipment makers across Asia, Europe, and North America. We do not bring a generic AI team — we bring people who know what a Cognex camera mount looks like and how an OPC-UA bus behaves under load.", False),
])
add_mixed_para([
    ("Reusable accelerators that cut time-to-pilot 40–60%. ", True),
    ("Vision-AI Starter Kit, Industrial IoT Reference Architecture, MES/SAP Integration Connector Library, MLOps Blueprint, and Edge Deployment Toolkit — field-proven, productized, licensed under the engagement. We are not building from zero on Day 1.", False),
])
add_mixed_para([
    ("Hyperscaler-aligned, integration-first architecture. ", True),
    ("AWS and Azure advanced partners. Pre-built integrations with SAP MES/ERP, SAP PM, AVEVA PI, OPC-UA, RFID, and major vision camera ecosystems. No rip-and-replace. Line-side latency below 200 ms.", False),
])
add_para("Reference outcomes from analogous engagements:", bold=True, size=10, color=NAVY, space_before=4)
add_bullet([("Tier-1 automotive OEM — paint defect CV. ", True),
            ("Scaled to 4 lines in 18 weeks. −58% escape defects, −42% rework cost, 14-month payback.", False)])
add_bullet([("Global EV manufacturer — predictive maintenance. ", True),
            ("Vibration + acoustic mesh on stamping & weld shops. −38% unplanned downtime, 5.2:1 ROI in year one.", False)])
add_bullet([("Aerospace Tier-1 — SOP pose-estimation. ", True),
            ("Multi-camera pipeline matched to SOP graph. 96% SOP compliance, audit prep time cut 70%.", False)])
add_bullet([("Heavy equipment — RFID + vision traceability. ", True),
            ("Unified Neo4j genealogy across 11 plants. Recall investigation cut from weeks to days.", False)])

# ============================================================================
# RECOMMENDATION
# ============================================================================
add_section_header("Recommendation: Two Pilots, 12 Weeks, Fixed Scope")
add_mixed_para([
    ("Pilot 1 — AI-based Visual Inspection at HMGMA Metaplant America. ", True),
    ("Edge-deployed CNN models on the IONIQ 5 / IONIQ 9 paint line. Success criteria: ≥30% reduction in escape defects vs baseline, <2s detection latency, MES line-hold integration validated. Fast-follow on Tucson / Santa Fe paint at HMMA.", False),
])
add_mixed_para([
    ("Pilot 2 — Predictive Maintenance on Ulsan stamping + Asan welding. ", True),
    ("Vibration + acoustic sensor mesh on 3–5 high-criticality assets at the world's largest auto plant. Success criteria: ≥25% reduction in unplanned downtime, work-order auto-generation in SAP PM, at least one predicted failure validated with measured lead time.", False),
])
add_mixed_para([
    ("Pilot team. ", True),
    ("PeopleTech embedded pod — solution architect, ML lead, vision engineer, DevOps, plant-floor integration specialist — co-located with a Hyundai owner team from week one.", False),
])
add_mixed_para([
    ("Gate to scale. ", True),
    ("Joint review at week 12. If both pilots clear KPIs, proceed to plant-wide rollout and queue the next two use cases (recommended: SOP Compliance + Variant Confirmation). If KPIs miss, Hyundai walks. No commitment to scale until data supports it.", False),
])
add_para(
    "Fixed-scope, fixed-fee engagement against an agreed pre-pilot baseline. The shape of the contract reflects the shape of the conviction.",
    italic=True, color=INK_SOFT, size=10
)

# ============================================================================
# NEXT STEP
# ============================================================================
add_section_header("Next Step")
add_mixed_para([
    ("Two-hour Pilot Scoping Workshop ", True),
    ("with Hyundai's Plant Operations and Quality leadership, plus PeopleTech's AI Manufacturing practice lead. We walk the candidate line, agree the baseline, finalize KPIs, and lock dates. ", False),
    ("Proposed timing: within the next 14 days. ", True),
    ("We hold a delivery pod open for Hyundai through the end of the quarter.", False),
])

# Final pull-quote
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PeopleTech is ready to build. Let's pick the first use case and move fast.")
r.font.name = "Calibri"; r.font.size = Pt(11); r.font.bold = True; r.font.italic = True
r.font.color.rgb = TEAL

doc.save("executive_brief.docx")
print("Saved: executive_brief.docx")
